import base64
import io
import unittest
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi.testclient import TestClient

import app
import engine
from resume_parser import extract_document, normalize, parse_sections


SAMPLE = '''张同学
教育背景
2021.09 - 2025.06 星河大学 计算机科学 本科
主修课程：数据结构、数据库
项目经历
项目名称：知识库问答平台
2024.03 - 2024.06 个人开发
负责文档处理和检索接口。
技术栈：Python、React，评估口径待确认。
项目名称：校园二手交易
2023.06 - 2023.09 团队项目
负责订单状态与幂等处理。
实习经历
2024.07 - 2024.09 星云科技 后端实习生
负责接口联调、日志排查和测试。
专业技能
Python、SQL
熟悉需求访谈和数据分析。
获奖经历
校级创新比赛一等奖
'''


class ResumeParserTests(unittest.TestCase):
    def test_four_sections_multiple_projects_short_lines_and_provenance(self):
        r=parse_sections(SAMPLE)
        self.assertEqual(r['summary']['项目'],2)
        self.assertEqual(r['summary']['实习经历'],1)
        self.assertEqual(r['summary']['学历'],1)
        self.assertEqual(r['summary']['工作技能'],1)
        self.assertIn('Python、SQL',r['sections']['工作技能'][0]['text'])
        self.assertIn('张同学',r['sections']['其他'][0]['text'])
        self.assertIn('一等奖',r['sections']['其他'][-1]['text'])
        self.assertTrue(all(c['source']['lines'] and not c['confirmed'] for c in r['cards']))
        self.assertIn('技术栈',r['sections']['项目'][0]['text'])
        self.assertNotIn('知识库',r['sections']['工作技能'][0]['text'])

    def test_heading_aliases_spacing_markdown_and_inline(self):
        r=parse_sections('## 教 育 背 景\n某某大学 本科\n【实习经历】：甲公司，负责接口测试。；项目：校园工具，设计检索。；工作技能：Python、用户访谈。')
        self.assertEqual([r['summary'][k] for k in ('学历','实习经历','项目','工作技能')],[1,1,1,1])
        en=parse_sections('EDUCATION\nExample University, Bachelor\nPROJECTS\nProject name: Catalog\nBuilt a catalog.\nINTERNSHIPS\nExample Inc, Intern\nTECHNICAL SKILLS\nSQL\nPython')
        self.assertEqual(en['summary']['学历'],1)
        self.assertEqual(en['summary']['实习经历'],1)
        self.assertIn('SQL',en['sections']['工作技能'][0]['text'])

    def test_no_truncation_after_eight_lines_or_long_entry(self):
        details=['负责功能'+str(i)+'，具体成果与指标需要验证。'*15 for i in range(120)]
        r=parse_sections('项目经历\n项目名称：大项目\n'+'\n'.join(details))
        joined='\n'.join(c['text'] for c in r['cards'])
        self.assertTrue(all(normalize(line) in joined for line in details))
        self.assertTrue(all(len(c['text'])<=12000 for c in r['cards']))

    def test_missing_sections_stay_empty_and_unknown_retained(self):
        r=parse_sections('张三\n自我评价\n做事细致，有责任心，愿意学习新领域。')
        for kind in ('学历','项目','实习经历','工作技能'):
            self.assertEqual(r['sections'][kind],[])
        self.assertTrue(r['sections']['其他'])

    def test_work_experience_not_mislabeled_as_internship(self):
        r=parse_sections('工作经历\n甲公司，正式员工，负责客户拓展。\n教育背景\n乙大学 本科')
        self.assertEqual(r['summary']['工作经历'],1)
        self.assertEqual(r['summary']['实习经历'],0)

    def test_project_title_before_second_date_stays_with_its_entry(self):
        r=parse_sections('项目经历\n校园商城\n2024.01 - 2024.03\n负责支付接口。\n知识库助手\n2024.05 - 2024.06\n负责文档导入。')
        self.assertEqual(r['summary']['项目'],2)
        self.assertNotIn('知识库助手',r['sections']['项目'][0]['text'])
        self.assertTrue(r['sections']['项目'][1]['text'].startswith('知识库助手'))

    def test_docx_layout_table_columns_and_textbox(self):
        from docx.oxml import parse_xml
        doc=Document();table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='教育背景';table.cell(1,0).text='甲大学 本科'
        table.cell(0,1).text='项目经历';table.cell(1,1).text='知识库平台，负责检索。'
        doc.element.body.append(parse_xml('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:txbxContent><w:p><w:r><w:t>工作技能：Python、SQL</w:t></w:r></w:p></w:txbxContent></w:r></w:p>'))
        stream=io.BytesIO();doc.save(stream)
        text,warnings=extract_document('layout.docx',stream.getvalue())
        r=parse_sections(text,warnings)
        self.assertIn('甲大学',r['sections']['学历'][0]['text'])
        self.assertNotIn('甲大学',r['sections']['项目'][0]['text'])
        self.assertIn('Python',r['sections']['工作技能'][0]['text'])

    def test_docx_tables_remain_in_document_order(self):
        doc=Document();doc.add_paragraph('教育背景')
        table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='某某大学';table.cell(0,1).text='本科 2021-2025'
        doc.add_paragraph('项目经历');doc.add_paragraph('项目名称：客服平台')
        doc.add_paragraph('负责客户问题分类和反馈分析。')
        doc.add_paragraph('工作技能');doc.add_paragraph('SQL、访谈')
        stream=io.BytesIO();doc.save(stream)
        text,w=extract_document('resume.docx',stream.getvalue())
        self.assertLess(text.index('某某大学'),text.index('项目经历'))
        result=parse_sections(text,w)
        self.assertIn('某某大学',result['sections']['学历'][0]['text'])
        self.assertNotIn('某某大学',result['sections']['工作技能'][0]['text'])

    def test_pdf_columns_and_missing_text_page(self):
        from reportlab.pdfgen.canvas import Canvas
        stream=io.BytesIO();c=Canvas(stream)
        # Draw rows interleaved to reproduce PDF content stream order issues.
        for y,left,right in [(790,'EDUCATION','PROJECTS'),(765,'Example University, Bachelor','Catalog application'),
                             (740,'2021 - 2025','Built search endpoints.'),(700,'SKILLS','INTERNSHIPS'),
                             (675,'Python, SQL','Example Inc - Intern'),(650,'User interviews','Built reporting tools.')]:
            c.drawString(40,y,left);c.drawString(330,y,right)
        c.showPage();c.showPage();c.save()
        text,warnings=extract_document('two-column.pdf',stream.getvalue())
        r=parse_sections(text,warnings)
        for kind in ('项目','实习经历','学历','工作技能'):
            self.assertEqual(r['summary'][kind],1,(kind,text,r['summary']))
        self.assertIn('Catalog',r['sections']['项目'][0]['text'])
        self.assertNotIn('Catalog',r['sections']['学历'][0]['text'])
        self.assertTrue(any('OCR' in w for w in warnings))

    def test_uploaded_file_and_text_share_schema(self):
        client=TestClient(app.app)
        txt=client.post('/api/resume',json={'text':SAMPLE})
        file=client.post('/api/resume',json={'name':'resume.txt','content':base64.b64encode(SAMPLE.encode()).decode()})
        self.assertEqual(txt.status_code,200)
        self.assertEqual(file.json()['sections'],txt.json()['sections'])
        self.assertEqual(file.json()['parser_version'],'sections-v2')

    def test_education_and_skills_are_not_projects_and_reach_llm(self):
        import asyncio
        r=parse_sections('教育背景\n星河大学 本科\n工作技能\n用户访谈与数据分析。')
        plan=engine.make_plan(r['cards'],{'name':'用户运营','design':'如何安排活动？'})
        self.assertFalse(any('项目深挖' in p['topic'] for p in plan))
        self.assertTrue(any('工作技能'==p['topic'] for p in plan))
        s={'plan':plan,'main_index':1,'depth':0,'messages':[],'mode':'live','jd':'分析留存',
           'profile':{'name':'运营'},'style':'严格','cards':r['cards']}
        mock=AsyncMock(return_value={'speaker':'interviewer','question':'请介绍你运用这些技能的经历。'})
        with patch.object(engine,'model_text',mock):
            asyncio.run(engine.next_question(s))
        payload=mock.call_args.args[1]
        self.assertEqual(payload['完整简历分类'],r['cards'])


if __name__=='__main__':unittest.main()
